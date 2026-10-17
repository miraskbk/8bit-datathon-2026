from __future__ import annotations

import argparse
import re
from pathlib import Path

import av
import numpy as np
import torch
from openpyxl import load_workbook
from torchvision.models.video import R3D_18_Weights, r3d_18


def parse_timestamp(value: str | None) -> float | None:
    if not value:
        return None
    minute = re.search(r"(\d+)\s*мин", value.lower())
    second = re.search(r"(\d+)\s*сек", value.lower())
    return float((int(minute.group(1)) * 60 if minute else 0) + (int(second.group(1)) if second else 0))


def read_labels(path: Path) -> dict[str, float | None]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    labels: dict[str, float | None] = {}
    for filename, timestamp, verdict in sheet.iter_rows(min_row=2, values_only=True):
        if not filename:
            continue
        labels[str(filename)] = parse_timestamp(timestamp) if verdict == "нарушение" else None
    return labels


def decode_at_four_fps(path: Path) -> tuple[torch.Tensor, np.ndarray]:
    frames = []
    timestamps = []
    with av.open(str(path)) as container:
        stream = container.streams.video[0]
        source_fps = float(stream.average_rate or 30)
        next_sample = 0.0
        for index, frame in enumerate(container.decode(stream)):
            timestamp = float(frame.time) if frame.time is not None else index / source_fps
            if timestamp + 1e-4 >= next_sample:
                array = np.array(frame.to_image().convert("RGB"), dtype=np.uint8, copy=True)
                frames.append(torch.from_numpy(array).permute(2, 0, 1))
                timestamps.append(timestamp)
                next_sample += 0.25
    return torch.stack(frames), np.asarray(timestamps)


def make_clips(
    frames: torch.Tensor,
    timestamps: np.ndarray,
    centers: list[float],
    transform,
) -> torch.Tensor:
    clips = []
    for center in centers:
        targets = np.linspace(center - 1.875, center + 1.875, 16)
        indices = [int(np.abs(timestamps - target).argmin()) for target in targets]
        clips.append(transform(frames[indices]))
    return torch.stack(clips)


def extract_features(
    labels_file: Path,
    videos_dir: Path,
    cache_path: Path,
) -> tuple[torch.Tensor, torch.Tensor]:
    labels = read_labels(labels_file)
    device = "cuda" if torch.cuda.is_available() else "cpu"
    weights = R3D_18_Weights.DEFAULT
    encoder = r3d_18(weights=weights)
    encoder.fc = torch.nn.Identity()
    encoder = encoder.to(device).eval()
    features = []
    targets = []
    for index, (stem, event_time) in enumerate(labels.items(), 1):
        path = videos_dir / f"{stem}.mov"
        if not path.is_file():
            raise FileNotFoundError(f"Video listed in labels was not found: {path}")
        frames, timestamps = decode_at_four_fps(path)
        duration = float(timestamps[-1])
        if event_time is not None:
            positive_centers = [
                max(2.0, min(duration - 2.0, event_time + offset))
                for offset in (-1, 0, 1)
            ]
            # Hard negatives from the same camera prevent the head from learning the
            # room/background instead of the action around the annotated timestamp.
            candidates = np.arange(2.0, max(2.01, duration - 1.75), 2.0)
            candidates = candidates[np.abs(candidates - event_time) >= 6.0]
            centers = positive_centers + candidates.tolist()
            clip_targets = [1.0] * len(positive_centers) + [0.0] * len(candidates)
        else:
            centers = np.arange(2.0, max(2.01, duration - 1.75), 2.0).tolist()
            clip_targets = [0.0] * len(centers)
        clips = make_clips(frames, timestamps, centers, weights.transforms())
        with torch.inference_mode():
            batches = [
                encoder(clips[start : start + 4].to(device)).cpu()
                for start in range(0, len(clips), 4)
            ]
        features.append(torch.cat(batches))
        targets.extend(clip_targets)
        print(f"[{index:02d}/{len(labels)}] {path.name}: {len(clips)} clips", flush=True)
    x = torch.cat(features)
    y = torch.tensor(targets, dtype=torch.float32)
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(cache_path, features=x.numpy(), labels=y.numpy())
    return x, y


def train_head(x: torch.Tensor, y: torch.Tensor, output: Path, epochs: int) -> None:
    torch.manual_seed(42)
    feature_mean = x.mean(dim=0)
    feature_std = x.std(dim=0).clamp_min(0.05)
    standardized = (x - feature_mean) / feature_std
    head = torch.nn.Linear(x.shape[1], 1)
    optimizer = torch.optim.AdamW(head.parameters(), lr=0.01, weight_decay=0.2)
    positive_weight = torch.tensor([(y == 0).sum() / (y == 1).sum()])
    loss_fn = torch.nn.BCEWithLogitsLoss(pos_weight=positive_weight)
    for epoch in range(epochs):
        optimizer.zero_grad()
        logits = head(standardized).squeeze(1)
        loss = loss_fn(logits, y)
        loss.backward()
        optimizer.step()
        if epoch % 100 == 0 or epoch == epochs - 1:
            accuracy = ((logits >= 0) == y.bool()).float().mean().item()
            print(f"epoch={epoch:03d} loss={loss.item():.4f} train_accuracy={accuracy:.3f}")
    output.parent.mkdir(parents=True, exist_ok=True)
    torch.save(
        {
            "encoder": "r3d_18_kinetics400_v1",
            "feature_mean": feature_mean,
            "feature_std": feature_std,
            "weight": head.weight.detach().squeeze(0),
            "bias": head.bias.detach().squeeze(0),
        },
        output,
    )
    print(f"Saved model head to {output}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Train the ForteBank temporal violation head")
    parser.add_argument("--labels", type=Path, default=Path("open/open_labels.xlsx"))
    parser.add_argument("--videos", type=Path, default=Path("open/videos"))
    parser.add_argument("--output", type=Path, default=Path("artifacts/forte_r3d_head.pt"))
    parser.add_argument("--feature-cache", type=Path, default=Path("artifacts/train_features.npz"))
    parser.add_argument("--reuse-features", action="store_true")
    parser.add_argument("--epochs", type=int, default=500)
    args = parser.parse_args()
    if args.reuse_features and args.feature_cache.is_file():
        cached = np.load(args.feature_cache)
        x = torch.tensor(cached["features"], dtype=torch.float32)
        y = torch.tensor(cached["labels"], dtype=torch.float32)
    else:
        x, y = extract_features(args.labels, args.videos, args.feature_cache)
    if y.unique().numel() != 2:
        raise ValueError("Both violation and no-violation videos are required")
    train_head(x, y, args.output, args.epochs)


def run() -> None:
    main()


if __name__ == "__main__":
    main()
